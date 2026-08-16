from __future__ import annotations

import hashlib
import html
import io
import json
import zipfile
from dataclasses import dataclass
from datetime import UTC, datetime

import matplotlib

matplotlib.use("Agg")

import matplotlib.pyplot as plt
import nbformat
import pandas as pd
from nbformat.validator import validate as validate_notebook
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from PIL import Image

from .config import APP_ID, APP_VERSION, CONTENT_VERSION, SCHEMA_VERSION


@dataclass(frozen=True)
class PackageBuildResult:
    filename: str
    data: bytes
    artifact_paths: list[str]


def vertical_slice_data() -> pd.DataFrame:
    return pd.DataFrame(
        [
            {"case_id": "CASE-001", "queue": "Billing", "case_age_hours": 12, "remaining_hours": 18},
            {"case_id": "CASE-002", "queue": "Billing", "case_age_hours": 28, "remaining_hours": 10},
            {"case_id": "CASE-003", "queue": "Technical", "case_age_hours": 8, "remaining_hours": 30},
            {"case_id": "CASE-004", "queue": "Technical", "case_age_hours": 35, "remaining_hours": 14},
            {"case_id": "CASE-005", "queue": "Retention", "case_age_hours": 20, "remaining_hours": 22},
            {"case_id": "CASE-006", "queue": "Retention", "case_age_hours": 42, "remaining_hours": 8},
        ]
    )


def _png_bytes(data: pd.DataFrame) -> bytes:
    summary = data.groupby("queue", sort=True)["remaining_hours"].mean()
    figure, axis = plt.subplots(figsize=(7.2, 4.2))
    colors = ["#2563EB", "#0F766E", "#B45309"]
    bars = axis.bar(summary.index, summary.values, color=colors)
    axis.set_title("Mean remaining hours by queue")
    axis.set_ylabel("Remaining hours")
    axis.set_xlabel("Queue")
    axis.spines[["top", "right"]].set_visible(False)
    axis.bar_label(bars, fmt="%.1f", padding=3)
    figure.tight_layout()
    buffer = io.BytesIO()
    figure.savefig(buffer, format="png", dpi=150, metadata={"Software": "ML Learning Lab"})
    plt.close(figure)
    return buffer.getvalue()


def _xlsx_bytes(data: pd.DataFrame) -> bytes:
    workbook = Workbook()
    overview = workbook.active
    overview.title = "Overview"
    overview.append(["ML Learning Lab — Milestone 1.1 evidence package"])
    overview.append(["Purpose", "Validate the end-to-end project export pipeline"])
    overview.append(["Rows", len(data)])
    overview.append(["Content version", CONTENT_VERSION])
    overview["A1"].font = Font(bold=True, size=14, color="FFFFFF")
    overview["A1"].fill = PatternFill("solid", fgColor="1D4ED8")
    overview.merge_cells("A1:B1")
    overview.column_dimensions["A"].width = 24
    overview.column_dimensions["B"].width = 58

    data_sheet = workbook.create_sheet("Sample Data")
    data_sheet.append(list(data.columns))
    for row in data.itertuples(index=False, name=None):
        data_sheet.append(list(row))
    for cell in data_sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="0F766E")
        cell.alignment = Alignment(horizontal="center")
    widths = {"A": 15, "B": 16, "C": 20, "D": 20}
    for column, width in widths.items():
        data_sheet.column_dimensions[column].width = width
    data_sheet.freeze_panes = "A2"

    dictionary = workbook.create_sheet("Data Dictionary")
    dictionary.append(["Field", "Meaning"])
    definitions = [
        ("case_id", "Synthetic case identifier"),
        ("queue", "Operational queue"),
        ("case_age_hours", "Age at the observation moment"),
        ("remaining_hours", "Fully observed time remaining until closure"),
    ]
    for definition in definitions:
        dictionary.append(definition)
    for cell in dictionary[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="7C3AED")
    dictionary.column_dimensions["A"].width = 22
    dictionary.column_dimensions["B"].width = 62

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _notebook_bytes() -> bytes:
    notebook = nbformat.v4.new_notebook(
        metadata={
            "kernelspec": {"display_name": "Python 3", "language": "python", "name": "python3"},
            "language_info": {"name": "python", "version": "3.13"},
        }
    )
    notebook.cells = [
        nbformat.v4.new_markdown_cell(
            "# Milestone 1.1 evidence analysis\n\nThis notebook uses the CSV included in the extracted package."
        ),
        nbformat.v4.new_code_cell(
            "from pathlib import Path\nimport pandas as pd\n\ndata = pd.read_csv(Path('data') / 'sample_cases.csv')\ndata.head()"
        ),
        nbformat.v4.new_code_cell(
            "queue_summary = data.groupby('queue', as_index=False)['remaining_hours'].mean()\nqueue_summary"
        ),
        nbformat.v4.new_markdown_cell(
            "The package is deliberately small: its purpose is to prove that learner work can leave the app as portable, readable artifacts."
        ),
    ]
    validate_notebook(notebook)
    return nbformat.writes(notebook).encode("utf-8")


def _briefs(data: pd.DataFrame) -> tuple[bytes, bytes]:
    queue = data.groupby("queue")["remaining_hours"].mean().idxmax()
    markdown = f"""# Milestone 1.1 evidence brief

## Question

Can the application produce a portable analyst package from deterministic synthetic case data?

## Result

Yes. The package contains {len(data)} synthetic cases. **{queue}** has the highest mean remaining time in this small validation sample.

## Limitation

This is pipeline evidence, not a business model or a completed course project.
"""
    body = "".join(
        f"<p>{html.escape(paragraph)}</p>"
        for paragraph in [
            "Can the application produce a portable analyst package from deterministic synthetic case data?",
            f"Yes. The package contains {len(data)} synthetic cases. {queue} has the highest mean remaining time in this small validation sample.",
            "This is pipeline evidence, not a business model or a completed course project.",
        ]
    )
    html_document = (
        "<!doctype html><html lang=\"en\"><head><meta charset=\"utf-8\">"
        "<title>Milestone 1.1 evidence brief</title></head><body>"
        "<h1>Milestone 1.1 evidence brief</h1>" + body + "</body></html>"
    )
    return markdown.encode("utf-8"), html_document.encode("utf-8")


def build_vertical_slice_package() -> PackageBuildResult:
    data = vertical_slice_data()
    markdown, html_document = _briefs(data)
    artifacts: dict[str, bytes] = {
        "data/sample_cases.csv": data.to_csv(index=False).encode("utf-8"),
        "charts/remaining_hours_by_queue.png": _png_bytes(data),
        "summary.xlsx": _xlsx_bytes(data),
        "analysis.ipynb": _notebook_bytes(),
        "brief.md": markdown,
        "brief.html": html_document,
    }
    manifest_entries = [
        {
            "path": path,
            "bytes": len(value),
            "sha256": hashlib.sha256(value).hexdigest(),
        }
        for path, value in sorted(artifacts.items())
    ]
    manifest_entries.append(
        {"path": "manifest.json", "bytes": None, "sha256": None, "role": "inventory"}
    )
    manifest = {
        "package_type": "milestone-1.1-vertical-slice",
        "schema_version": SCHEMA_VERSION,
        "app_id": APP_ID,
        "content_version": CONTENT_VERSION,
        "app_version": APP_VERSION,
        "generated_at": datetime.now(UTC).isoformat(),
        "artifacts": manifest_entries,
    }
    artifacts["manifest.json"] = json.dumps(manifest, indent=2).encode("utf-8")

    package = io.BytesIO()
    with zipfile.ZipFile(package, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for path, value in sorted(artifacts.items()):
            archive.writestr(path, value)
    return PackageBuildResult(
        filename="ml-learning-lab-milestone1.1-evidence.zip",
        data=package.getvalue(),
        artifact_paths=sorted(artifacts),
    )


def validate_package(package_data: bytes) -> dict[str, object]:
    required = {
        "data/sample_cases.csv",
        "charts/remaining_hours_by_queue.png",
        "summary.xlsx",
        "analysis.ipynb",
        "brief.md",
        "brief.html",
        "manifest.json",
    }
    with zipfile.ZipFile(io.BytesIO(package_data)) as archive:
        bad_file = archive.testzip()
        if bad_file:
            raise ValueError(f"ZIP CRC failed for {bad_file}")
        names = set(archive.namelist())
        if names != required:
            raise ValueError(f"Unexpected package inventory: {sorted(names ^ required)}")
        data = pd.read_csv(archive.open("data/sample_cases.csv"))
        if data.empty or set(data.columns) != {
            "case_id",
            "queue",
            "case_age_hours",
            "remaining_hours",
        }:
            raise ValueError("CSV data is empty or has the wrong schema")
        with Image.open(io.BytesIO(archive.read("charts/remaining_hours_by_queue.png"))) as image:
            image.verify()
        workbook = load_workbook(io.BytesIO(archive.read("summary.xlsx")), read_only=True)
        if workbook.sheetnames != ["Overview", "Sample Data", "Data Dictionary"]:
            raise ValueError("Workbook sheets do not match the package contract")
        notebook_text = archive.read("analysis.ipynb").decode("utf-8")
        notebook = nbformat.reads(notebook_text, as_version=4)
        validate_notebook(notebook)
        if "data') / 'sample_cases.csv" not in notebook_text:
            raise ValueError("Notebook does not use the included relative CSV path")
        markdown = archive.read("brief.md").decode("utf-8")
        html_text = archive.read("brief.html").decode("utf-8")
        if not markdown.startswith("# ") or "<!doctype html>" not in html_text.lower():
            raise ValueError("Markdown or HTML brief is invalid")
        manifest = json.loads(archive.read("manifest.json").decode("utf-8"))
        if manifest.get("package_type") != "milestone-1.1-vertical-slice":
            raise ValueError("Manifest package type does not match Milestone 1.1")
        if manifest.get("app_id") != APP_ID or manifest.get("app_version") != APP_VERSION:
            raise ValueError("Manifest application identity does not match this build")
        inventory = {entry["path"] for entry in manifest["artifacts"]}
        if inventory != required:
            raise ValueError("Manifest does not inventory every artifact")
        for entry in manifest["artifacts"]:
            if entry["path"] == "manifest.json":
                continue
            value = archive.read(entry["path"])
            if hashlib.sha256(value).hexdigest() != entry["sha256"]:
                raise ValueError(f"Manifest hash failed for {entry['path']}")
    return {
        "valid": True,
        "artifact_count": len(required),
        "rows": len(data),
        "artifacts": sorted(required),
    }
